import os
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

def create_staynest_excel():
    excel_path = "excel/StayNest_Hotel_Booking_Analysis.xlsx"
    os.makedirs(os.path.dirname(excel_path), exist_ok=True)
    df = pd.read_csv("cleaned_data/staynest_cleaned_data.csv")
    
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    navy_dark = "1B365D"
    blue_header = "2980B9"
    soft_blue = "D4E6F1"
    light_gray = "F8F9F9"
    border_gray = "BDC3C7"
    
    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=13, bold=True, color=navy_dark)
    font_tbl_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="000000")
    font_regular = Font(name="Calibri", size=11, color="000000")
    
    fill_dark_header = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
    fill_blue_header = PatternFill(start_color=blue_header, end_color=blue_header, fill_type="solid")
    fill_kpi_bg = PatternFill(start_color=soft_blue, end_color=soft_blue, fill_type="solid")
    fill_zebra = PatternFill(start_color=light_gray, end_color=light_gray, fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color=border_gray)
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # SHEET 1: EXECUTIVE SUMMARY
    ws_sum = wb.create_sheet(title="Executive Summary")
    ws_sum.views.sheetView[0].showGridLines = True
    
    ws_sum.merge_cells("B2:J2")
    ws_sum["B2"] = "STAYNEST HOSPITALITY GROUP - HOTEL BOOKING INTELLIGENCE"
    ws_sum["B2"].font = font_title
    ws_sum["B2"].fill = fill_dark_header
    ws_sum["B2"].alignment = align_center
    ws_sum.row_dimensions[2].height = 35
    
    ws_sum.merge_cells("B3:J3")
    ws_sum["B3"] = "DS Sprint - Skill Based Assessment 2 | Use Case 10: Understand Bookings, Reduce Cancellations, Improve Occupancy"
    ws_sum["B3"].font = font_subtitle
    ws_sum["B3"].fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
    ws_sum["B3"].alignment = align_center
    ws_sum.row_dimensions[3].height = 22
    
    kpis = [
        ("TOTAL BOOKINGS", f"{len(df):,}", "B5:C6"),
        ("OVERALL CANCEL RATE", f"{df['is_canceled'].mean()*100:.1f}%", "D5:D6"),
        ("REALIZED REVENUE", f"${df[df['is_canceled']==0]['revenue'].sum()/1e6:.2f}M", "E5:F6"),
        ("AVERAGE DAILY RATE", f"${df['adr'].mean():.2f}", "G5:G6"),
        ("AVG LEAD TIME", f"{df['lead_time'].mean():.1f} Days", "H5:H6"),
        ("REPEAT GUEST RETENTION", f"{df['is_repeated_guest'].mean()*100:.1f}%", "I5:J6")
    ]
    
    for label, val, cell_range in kpis:
        ws_sum.merge_cells(cell_range)
        top_cell = cell_range.split(":")[0]
        ws_sum[top_cell] = f"{val}\n{label}"
        ws_sum[top_cell].font = Font(name="Calibri", size=13, bold=True, color=navy_dark)
        ws_sum[top_cell].fill = fill_kpi_bg
        ws_sum[top_cell].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws_sum[cell_range]:
            for c in row:
                c.border = thin_border
    ws_sum.row_dimensions[5].height = 26
    ws_sum.row_dimensions[6].height = 26
    
    ws_sum["B8"] = "CORE ANALYTICAL FINDINGS & ACTION MATRIX"
    ws_sum["B8"].font = font_section
    
    summary_headers = ["Analysis Area", "Key Observation / Evidence", "Metric Benchmark", "Management Action / Recommendation"]
    ws_sum.merge_cells("C9:E9")
    ws_sum.merge_cells("F9:G9")
    ws_sum.merge_cells("H9:J9")
    ws_sum["B9"] = summary_headers[0]
    ws_sum["C9"] = summary_headers[1]
    ws_sum["F9"] = summary_headers[2]
    ws_sum["H9"] = summary_headers[3]
    
    for c_idx in ["B", "C", "D", "E", "F", "G", "H", "I", "J"]:
        ws_sum[f"{c_idx}9"].fill = fill_blue_header
        ws_sum[f"{c_idx}9"].font = font_tbl_header
        ws_sum[f"{c_idx}9"].alignment = align_center
        ws_sum[f"{c_idx}9"].border = thin_border
        
    summary_rows = [
        ("1. Channel Cancellation", "TA/TO intermediary channel has highest cancellations; Direct/Corporate show high loyalty.", "TA/TO 31.16% vs Direct 14.98%", "Shift marketing budget toward Direct & Loyalty programs; impose non-refundable OTA policies."),
        ("2. Lead Time Risk", "Cancellations accelerate heavily on bookings made >90 days out.", "0-14d (11.63%) vs 180d+ (39.85%)", "Implement stepped non-refundable advance deposits on bookings placed >60 days in advance."),
        ("3. Seasonal Occupancy", "Summer delivers peak volume and ADR; Winter experiences demand drop.", "Summer $16.60M vs Winter $3.96M", "Deploy peak yield dynamic pricing in Summer; launch corporate and winter retreat packages."),
        ("4. Repeat Guests", "Repeat guests are 3.45x less likely to cancel than first-time visitors.", "Repeat 8.24% vs First-Time 28.42%", "Formalize StayNest Perks program to boost repeat visitation and protect base occupancy."),
        ("5. Room Mix", "Room Type A is the core revenue volume driver; Type D & Suites deliver high ADR.", "Room A $11.98M (93 ADR); Room D $5.69M", "Maintain core pricing on Room A; use automated upgrade nudges to upsell Rooms D, E & Suites."),
        ("6. Event Impact", "Major events and cultural festivals unlock strong pricing power.", "ADR $125.59 (Events) vs $102.83 (Normal)", "Implement festival surge pricing and enforce 3-night Minimum Length of Stay (MLOS).")
    ]
    
    curr_r = 10
    for area, obs, bench, act in summary_rows:
        ws_sum[f"B{curr_r}"] = area
        ws_sum.merge_cells(f"C{curr_r}:E{curr_r}")
        ws_sum[f"C{curr_r}"] = obs
        ws_sum.merge_cells(f"F{curr_r}:G{curr_r}")
        ws_sum[f"F{curr_r}"] = bench
        ws_sum.merge_cells(f"H{curr_r}:J{curr_r}")
        ws_sum[f"H{curr_r}"] = act
        ws_sum[f"B{curr_r}"].font = font_bold
        ws_sum[f"C{curr_r}"].font = font_regular
        ws_sum[f"F{curr_r}"].font = font_bold
        ws_sum[f"H{curr_r}"].font = font_regular
        ws_sum[f"B{curr_r}"].alignment = align_left
        ws_sum[f"C{curr_r}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws_sum[f"F{curr_r}"].alignment = align_center
        ws_sum[f"H{curr_r}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        fill = fill_zebra if curr_r % 2 == 0 else PatternFill(fill_type=None)
        for c_idx in ["B", "C", "D", "E", "F", "G", "H", "I", "J"]:
            ws_sum[f"{c_idx}{curr_r}"].border = thin_border
            if fill.fill_type:
                ws_sum[f"{c_idx}{curr_r}"].fill = fill
        ws_sum.row_dimensions[curr_r].height = 36
        curr_r += 1
        
    for col in ["B", "C", "D", "E", "F", "G", "H", "I", "J"]:
        ws_sum.column_dimensions[col].width = 18

    # HELPER FOR SHEETS WITH CHARTS
    def create_sheet_with_chart(title, headers, rows, chart_col, chart_title, is_pct=True):
        ws = wb.create_sheet(title=title)
        ws.views.sheetView[0].showGridLines = True
        ws.merge_cells("B2:G2")
        ws["B2"] = f"STAYNEST HOSPITALITY - {title.upper()}"
        ws["B2"].font = font_title
        ws["B2"].fill = fill_dark_header
        ws["B2"].alignment = align_center
        ws.row_dimensions[2].height = 30
        
        ws.row_dimensions[4].height = 24
        for c_idx, h in enumerate(headers, start=2):
            cell = ws.cell(row=4, column=c_idx, value=h)
            cell.font = font_tbl_header
            cell.fill = fill_blue_header
            cell.alignment = align_center
            cell.border = thin_border
            
        for r_idx, r_vals in enumerate(rows, start=5):
            fill = fill_zebra if r_idx % 2 == 0 else PatternFill(fill_type=None)
            for c_idx, val in enumerate(r_vals, start=2):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = font_regular
                cell.border = thin_border
                if fill.fill_type:
                    cell.fill = fill
                if isinstance(val, (int, np.integer)):
                    cell.number_format = '#,##0'
                    cell.alignment = align_right
                elif isinstance(val, (float, np.floating)):
                    cell.number_format = '0.00"%"' if (c_idx == chart_col and is_pct) else ('$#,##0.00' if "ADR" in headers[c_idx-2] or "Revenue" in headers[c_idx-2] else '#,##0.00')
                    cell.alignment = align_right
                else:
                    cell.alignment = align_left
                    
        chart = BarChart()
        chart.type = "col"
        chart.title = chart_title
        chart.y_axis.title = "% Rate" if is_pct else "Volume"
        data_ref = Reference(ws, min_col=chart_col, min_row=4, max_row=4+len(rows))
        cats_ref = Reference(ws, min_col=2, min_row=5, max_row=4+len(rows))
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.legend = None
        chart.width = 16
        chart.height = 10
        ws.add_chart(chart, f"B{len(rows) + 8}")
        
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # 1. Channel Cancellation
    ch_df = df.groupby('booking_channel').agg(
        total=('is_canceled', 'count'),
        canceled=('is_canceled', 'sum'),
        completed=('is_canceled', lambda x: (x == 0).sum()),
        rate=('is_canceled', lambda x: round(x.mean() * 100, 2)),
        adr=('adr', lambda x: round(x.mean(), 2))
    ).reset_index().sort_values('rate', ascending=False)
    create_sheet_with_chart("Channel Cancellation", ["Booking Channel", "Total Bookings", "Cancellations", "Completed Stays", "Cancellation Rate (%)", "Average ADR ($)"], [list(r) for r in ch_df.values], 6, "Cancellation Rate (%) by Channel")

    # 2. Lead Time Analysis
    lt_order = ['0-14 days (Short)', '15-30 days (Moderate)', '31-90 days (Quarterly)', '91-180 days (Long)', '180+ days (Very Long)']
    lt_df = df.groupby('lead_time_group').agg(
        total=('is_canceled', 'count'),
        canceled=('is_canceled', 'sum'),
        completed=('is_canceled', lambda x: (x == 0).sum()),
        rate=('is_canceled', lambda x: round(x.mean() * 100, 2)),
        avg_adr=('adr', lambda x: round(x.mean(), 2))
    ).reindex(lt_order).reset_index()
    create_sheet_with_chart("Lead Time Analysis", ["Lead Time Bracket", "Total Bookings", "Cancellations", "Completed Stays", "Cancellation Rate (%)", "Average ADR ($)"], [list(r) for r in lt_df.values], 6, "Cancellation Rate (%) Across Lead Time Brackets")

    # 3. Seasonal Occupancy
    s_df = df.groupby('season').agg(
        total=('is_canceled', 'count'),
        completed=('is_canceled', lambda x: (x == 0).sum()),
        canceled=('is_canceled', 'sum'),
        rate=('is_canceled', lambda x: round(x.mean() * 100, 2)),
        adr=('adr', lambda x: round(x.mean(), 2)),
        revenue=('revenue', lambda x: round(x[df.loc[x.index, 'is_canceled'] == 0].sum(), 2))
    ).reindex(['Summer', 'Spring', 'Autumn', 'Winter']).reset_index()
    create_sheet_with_chart("Seasonal Occupancy", ["Season", "Total Bookings", "Completed Stays", "Cancellations", "Cancellation Rate (%)", "Average ADR ($)", "Realized Revenue ($)"], [list(r) for r in s_df.values], 8, "Realized Room Revenue ($) by Season", is_pct=False)

    # 4. Guest Type Comparison
    g_df = df.groupby('repeat_status').agg(
        total=('is_canceled', 'count'),
        completed=('is_canceled', lambda x: (x == 0).sum()),
        canceled=('is_canceled', 'sum'),
        rate=('is_canceled', lambda x: round(x.mean() * 100, 2)),
        adr=('adr', lambda x: round(x.mean(), 2))
    ).reset_index()
    create_sheet_with_chart("Guest-Type Comparison", ["Guest Loyalty", "Total Bookings", "Completed Stays", "Cancellations", "Cancellation Rate (%)", "Average ADR ($)"], [list(r) for r in g_df.values], 6, "Cancellation Rate (%) by Loyalty Status")

    # 5. Room Performance
    r_df = df.groupby('room_type').agg(
        total=('is_canceled', 'count'),
        completed=('is_canceled', lambda x: (x == 0).sum()),
        canceled=('is_canceled', 'sum'),
        rate=('is_canceled', lambda x: round(x.mean() * 100, 2)),
        adr=('adr', lambda x: round(x.mean(), 2)),
        revenue=('revenue', lambda x: round(x[df.loc[x.index, 'is_canceled'] == 0].sum(), 2))
    ).reset_index().sort_values('revenue', ascending=False)
    create_sheet_with_chart("Room Performance", ["Room Type", "Total Bookings", "Completed Stays", "Cancellations", "Cancellation Rate (%)", "Average ADR ($)", "Realized Revenue ($)"], [list(r) for r in r_df.values], 8, "Realized Revenue ($) by Room Category", is_pct=False)

    # 6. Event Impact
    e_df = df.groupby('has_event').agg(
        total=('is_canceled', 'count'),
        completed=('is_canceled', lambda x: (x == 0).sum()),
        canceled=('is_canceled', 'sum'),
        rate=('is_canceled', lambda x: round(x.mean() * 100, 2)),
        adr=('adr', lambda x: round(x.mean(), 2)),
        revenue=('revenue', lambda x: round(x[df.loc[x.index, 'is_canceled'] == 0].sum(), 2))
    ).reset_index()
    e_df['Day Type'] = e_df['has_event'].map({0: 'Regular Non-Event Days', 1: 'Event / Festival Days'})
    e_df = e_df[['Day Type', 'total', 'completed', 'canceled', 'rate', 'adr', 'revenue']]
    create_sheet_with_chart("Event Impact", ["Day Category", "Total Bookings", "Completed Stays", "Cancellations", "Cancellation Rate (%)", "Average ADR ($)", "Realized Revenue ($)"], [list(r) for r in e_df.values], 7, "Average ADR ($): Event vs Normal Days", is_pct=False)

    wb.save(excel_path)
    print(f"Generated: {excel_path}")

if __name__ == "__main__":
    create_staynest_excel()
